from sqlalchemy.orm import Session
from sqlalchemy import func
from ..models.user import User, Grade, Class, TeacherClass
from ..models.task import Task, AnswerSheet, AnswerRecord
from ..models.risk import RiskAlert, Intervention, ScoringResult, QualityAssessment
from ..schemas.user import UserCreate, UserUpdate, UserInfo, ImportResult
from ..utils.password import hash_password
from ..utils.validators import validate_id_card, mask_phone
import openpyxl
from io import BytesIO


def list_users(db: Session, school_id: int, role: str, page: int = 1, page_size: int = 20,
               grade_id: int | None = None, class_id: int | None = None, keyword: str = "",
               teacher_type: str = "", class_ids: list[int] | None = None):
    q = db.query(User).filter(User.school_id == school_id, User.role == role)
    if grade_id:
        q = q.filter(User.grade_id == grade_id)
    if class_id:
        q = q.filter(User.class_id == class_id)
    if class_ids is not None:
        q = q.filter(User.class_id.in_(class_ids))
    if keyword:
        safe_kw = keyword.replace("%", "\\%").replace("_", "\\_")
        q = q.filter(User.real_name.contains(safe_kw) | User.username.contains(safe_kw) | User.student_no.contains(safe_kw))
    if teacher_type and role in ("teacher", "counselor"):
        q = q.filter(User.teacher_type == teacher_type)

    total = q.count()
    items = q.order_by(User.id.desc()).offset((page - 1) * page_size).limit(page_size).all()

    result = []
    for u in items:
        grade_name = db.query(Grade).filter(Grade.id == u.grade_id).first()
        class_name = db.query(Class).filter(Class.id == u.class_id).first()
        result.append(
            UserInfo(
                id=u.id, school_id=u.school_id, username=u.username, real_name=u.real_name,
                role=u.role, teacher_type=u.teacher_type, gender=u.gender or "", phone=mask_phone(u.phone),
                student_no=u.student_no or "", birth_date=u.birth_date, grade_id=u.grade_id,
                class_id=u.class_id, status=u.status,
                grade_name=grade_name.name if grade_name else "",
                class_name=class_name.name if class_name else "",
                created_at=u.created_at,
            )
        )

    return {"items": result, "total": total, "page": page, "page_size": page_size, "total_pages": max((total + page_size - 1) // page_size, 1)}


def list_teachers(db: Session, school_id: int, page: int = 1, page_size: int = 20,
                  keyword: str = "", teacher_type: str = ""):
    q = db.query(User).filter(User.school_id == school_id, User.role.in_(["teacher", "counselor"]))
    if keyword:
        safe_kw = keyword.replace("%", "\\%").replace("_", "\\_")
        q = q.filter(User.real_name.contains(safe_kw) | User.username.contains(safe_kw))
    if teacher_type:
        q = q.filter(User.teacher_type == teacher_type)

    total = q.count()
    items = q.order_by(User.id.desc()).offset((page - 1) * page_size).limit(page_size).all()

    result = []
    for u in items:
        grade_name = db.query(Grade).filter(Grade.id == u.grade_id).first()
        class_name = db.query(Class).filter(Class.id == u.class_id).first()
        result.append(
            UserInfo(
                id=u.id, school_id=u.school_id, username=u.username, real_name=u.real_name,
                role=u.role, teacher_type=u.teacher_type, gender=u.gender or "", phone=mask_phone(u.phone),
                student_no=u.student_no or "", birth_date=u.birth_date, grade_id=u.grade_id,
                class_id=u.class_id, status=u.status,
                grade_name=grade_name.name if grade_name else "",
                class_name=class_name.name if class_name else "",
                created_at=u.created_at,
            )
        )

    return {"items": result, "total": total, "page": page, "page_size": page_size, "total_pages": max((total + page_size - 1) // page_size, 1)}


def _generate_student_no(db: Session, school_id: int) -> str:
    """自动生成唯一学号: STU{school_id:03d}{seq:04d}"""
    prefix = f"STU{school_id:03d}"
    max_no = db.query(func.max(User.student_no)).filter(
        User.school_id == school_id, User.role == "student",
        User.student_no.like(f"{prefix}%")
    ).scalar()
    if max_no and len(max_no) >= len(prefix) + 4:
        seq = int(max_no[-4:]) + 1
    else:
        seq = 1
    return f"{prefix}{seq:04d}"


def create_user(db: Session, data: UserCreate) -> User:
    role = data.role or "student"
    # 学生的用户名是身份证号，需要校验并规范化（小写x→大写X）
    username = data.username
    if role == "student":
        username = validate_id_card(username)
    default_pw = "123456"
    force_change = False
    if role in ("student", "teacher", "counselor") and len(username) >= 6:
        default_pw = username[-6:]
        force_change = True
    student_no = data.student_no
    if role == "student" and not student_no:
        student_no = _generate_student_no(db, data.school_id)
    u = User(
        school_id=data.school_id, username=username,
        password_hash=hash_password(data.password or default_pw),
        real_name=data.real_name, role=role, teacher_type=data.teacher_type,
        gender=data.gender, phone=data.phone, student_no=student_no or "",
        birth_date=data.birth_date, grade_id=data.grade_id, class_id=data.class_id, status=data.status,
        must_change_password=force_change,
    )
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def update_user(db: Session, user_id: int, data: UserUpdate) -> User:
    u = db.query(User).filter(User.id == user_id).first()
    if not u:
        raise ValueError("用户不存在")
    update_data = data.model_dump(exclude_unset=True)
    # 学生身份证号更新时规范化（小写x→大写X）
    if "username" in update_data and u.role == "student":
        update_data["username"] = validate_id_card(update_data["username"])
    for k, v in update_data.items():
        setattr(u, k, v)
    db.commit()
    db.refresh(u)
    return u


def delete_user(db: Session, user_id: int):
    u = db.query(User).filter(User.id == user_id).first()
    if not u:
        return
    if u.role == "student":
        # 学生：级联删除答卷、评分、质检、风险预警、干预
        sheet_ids = [s.id for s in db.query(AnswerSheet.id).filter(AnswerSheet.student_id == user_id).all()]
        if sheet_ids:
            db.query(AnswerRecord).filter(AnswerRecord.answer_sheet_id.in_(sheet_ids)).delete(synchronize_session=False)
            db.query(ScoringResult).filter(ScoringResult.answer_sheet_id.in_(sheet_ids)).delete(synchronize_session=False)
            db.query(QualityAssessment).filter(QualityAssessment.answer_sheet_id.in_(sheet_ids)).delete(synchronize_session=False)
            db.query(RiskAlert).filter(RiskAlert.answer_sheet_id.in_(sheet_ids)).delete(synchronize_session=False)
            db.query(AnswerSheet).filter(AnswerSheet.id.in_(sheet_ids)).delete(synchronize_session=False)
        db.query(RiskAlert).filter(RiskAlert.student_id == user_id).delete()
        db.query(Intervention).filter(Intervention.student_id == user_id).delete()
    elif u.role in ("teacher", "counselor"):
        # 教师：级联删除教师-班级关联、干预记录
        db.query(TeacherClass).filter(TeacherClass.teacher_id == user_id).delete()
        db.query(Intervention).filter(Intervention.teacher_id == user_id).delete()
        db.query(Class).filter(Class.head_teacher_id == user_id).update({Class.head_teacher_id: None})
        db.query(Class).filter(Class.counselor_id == user_id).update({Class.counselor_id: None})
        db.query(RiskAlert).filter(RiskAlert.assigned_teacher_id == user_id).update({RiskAlert.assigned_teacher_id: None})
    db.delete(u)
    db.commit()


def get_user_info(db: Session, user_id: int) -> UserInfo:
    u = db.query(User).filter(User.id == user_id).first()
    if not u:
        raise ValueError("用户不存在")
    grade_name = db.query(Grade).filter(Grade.id == u.grade_id).first()
    class_name = db.query(Class).filter(Class.id == u.class_id).first()
    return UserInfo(
        id=u.id, school_id=u.school_id, username=u.username, real_name=u.real_name,
        role=u.role, teacher_type=u.teacher_type, gender=u.gender or "", phone=mask_phone(u.phone),
        student_no=u.student_no or "", birth_date=u.birth_date, grade_id=u.grade_id,
        class_id=u.class_id, status=u.status,
        grade_name=grade_name.name if grade_name else "",
        class_name=class_name.name if class_name else "",
        created_at=u.created_at,
    )


def import_students_from_excel(db: Session, school_id: int, file_bytes: bytes) -> ImportResult:
    result = ImportResult()
    wb = openpyxl.load_workbook(BytesIO(file_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # Read header row to detect column mapping
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    col_map = {}
    for idx, h in enumerate(header_row):
        if h:
            col_map[str(h).strip()] = idx

    grades_map = {g.name: g.id for g in db.query(Grade).filter(Grade.school_id == school_id).all()}
    classes_map = {c.name: c.id for c in db.query(Class).filter(Class.school_id == school_id).all()}

    def get_val(row_tuple, col_name, fallback_idx=None):
        if col_name in col_map:
            return row_tuple[col_map[col_name]]
        if fallback_idx is not None and fallback_idx < len(row_tuple):
            return row_tuple[fallback_idx]
        return None

    for i, row in enumerate(rows):
        try:
            student_no = str(get_val(row, "学号") or "").strip()
            name = str(get_val(row, "姓名", 1) or "").strip()
            id_card = str(get_val(row, "身份证号") or "").strip() if "身份证号" in col_map else ""
            gender = str(get_val(row, "性别", 2) or "").strip()
            grade_name = str(get_val(row, "年级", 3) or "").strip()
            class_name = str(get_val(row, "班级", 4) or "").strip()
            birth_date_str = get_val(row, "出生日期", 5)
            phone = str(get_val(row, "手机号", 6) or "").strip()

            if not name:
                result.fail_count += 1
                result.errors.append(f"第{i + 2}行: 姓名为空")
                continue

            if not id_card:
                result.fail_count += 1
                result.errors.append(f"第{i + 2}行: 身份证号不能为空")
                continue

            username = id_card
            if id_card:
                try:
                    username = validate_id_card(id_card)  # 使用校验后的规范化值（大写X）
                except ValueError as ve:
                    result.fail_count += 1
                    result.errors.append(f"第{i + 2}行: {str(ve)}")
                    continue

            grade_id = grades_map.get(grade_name)
            class_id = classes_map.get(class_name)

            existing = db.query(User).filter(User.username == username).first()
            if existing:
                result.fail_count += 1
                result.errors.append(f"第{i + 2}行: 身份证号 {username[:3]}****{username[-4:]} 已存在")
                continue

            from datetime import datetime
            birth_date = None
            if birth_date_str:
                try:
                    if isinstance(birth_date_str, datetime):
                        birth_date = birth_date_str.date()
                    else:
                        birth_date = datetime.strptime(str(birth_date_str).strip(), "%Y-%m-%d").date()
                except ValueError:
                    pass

            # 自动生成学号（如果导入文件中没有提供）
            if not student_no:
                student_no = _generate_student_no(db, school_id)
            default_pw = username[-6:] if len(username) >= 6 else "123456"
            u = User(
                school_id=school_id, username=username, password_hash=hash_password(default_pw),
                real_name=name, role="student", gender=gender, student_no=student_no,
                phone=phone, birth_date=birth_date, grade_id=grade_id, class_id=class_id,
                status=True, must_change_password=True,
            )
            db.add(u)
            result.success_count += 1
        except Exception as e:
            result.fail_count += 1
            result.errors.append(f"第{i + 2}行: {str(e)}")

    db.commit()
    return result


def generate_student_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学生导入模板"
    ws.append(["姓名", "身份证号", "性别", "年级", "班级", "出生日期", "手机号"])
    ws.append(["张三", "610102201001150015", "男", "初一", "初一(1)班", "2010-01-15", "13800000000"])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def assign_teacher_classes(db: Session, teacher_id: int, class_ids: list[int]):
    db.query(TeacherClass).filter(TeacherClass.teacher_id == teacher_id).delete()
    for cid in class_ids:
        db.add(TeacherClass(teacher_id=teacher_id, class_id=cid))
    db.commit()
